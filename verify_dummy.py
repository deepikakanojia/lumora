# Replays getDashboardData() from the onboarding Apps Script against the generated workbook.
import openpyxl, datetime as dt, re, sys
from demo_config import REAL_DOMAIN, real_domain_pattern
from openpyxl.utils import get_column_letter as L
wb=openpyxl.load_workbook("NHO_Tracker_Dummy_Data.xlsx")
print("tabs:", [w.title for w in wb.worksheets])
print("tab0 len:", len(wb.worksheets[0].title))
ws=wb.worksheets[0]
print("freeze:", ws.freeze_panes, "| header row 7 col E:", repr(ws.cell(7,5).value))

COHORT_DATES=["2026-05-11","2026-05-25","2026-06-08","2026-06-22","2026-07-06","2026-07-20",
 "2026-08-03","2026-08-17","2026-08-31","2026-09-14","2026-09-28","2026-10-12","2026-10-26",
 "2026-11-09","2026-11-23","2026-12-07","2026-12-21"]
fmt=lambda d: d.strftime("%b %d, %Y")           # Utilities.formatDate 'MMM dd, yyyy'
data={d:[] for d in COHORT_DATES}
leak=[]
# Pattern comes from demo_config; None when no real domain is configured.
PII_DOMAIN=real_domain_pattern()
_LABEL=REAL_DOMAIN or '<no real domain configured>'
for r in range(2, ws.max_row+1):
    fn=ws.cell(r,5).value; ln=ws.cell(r,6).value
    if not (fn or ln): continue
    h=ws.cell(r,8).value
    cohort_cell = fmt(h) if isinstance(h,dt.datetime) else str(h).strip()
    for d in COHORT_DATES:
        if cohort_cell == fmt(dt.datetime.strptime(d,"%Y-%m-%d")):
            data[d].append(f"{fn} {ln}")
    for c in range(1,26):
        v=ws.cell(r,c).value
        if PII_DOMAIN and isinstance(v,str) and PII_DOMAIN.search(v): leak.append(f"{L(c)}{r}")

matched=sum(len(v) for v in data.values())
print(f"\ndashboard matched {matched} hires into cohorts")
for d in COHORT_DATES:
    print(f"  {fmt(dt.datetime.strptime(d,'%Y-%m-%d'))}: {len(data[d])} hire(s)")
unmatched = 48-matched
print("\nunmatched hires (would be invisible on dashboard):", unmatched)
print(f"{_LABEL} leaks:", leak or ("none" if PII_DOMAIN else "SKIPPED - set DEMO_REAL_DOMAIN"))

# PII scan across every tab
import string
bad=[]
for w in wb.worksheets:
    for row in w.iter_rows():
        for cell in row:
            if PII_DOMAIN and isinstance(cell.value,str) and PII_DOMAIN.search(cell.value):
                bad.append((w.title,cell.coordinate))
print(f"{_LABEL} anywhere in workbook:", bad or ("none" if PII_DOMAIN else "SKIPPED - set DEMO_REAL_DOMAIN"))
sys.exit(0 if unmatched==0 and not bad else 1)
