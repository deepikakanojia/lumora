# Generates NHO_Dummy_Extensions.xlsx — the synthetic companion workbook the
# redesigned dashboard needs but the tracker clone does not carry.
#
# Why a SECOND workbook rather than extra columns on the first:
#   NHO_Tracker_Dummy_Data.xlsx is a clone of the live
#   "Cohort Tracker May - December 26" tab, whose columns are read by position.
#   Inserting anything into it would break the very contract it exists to
#   demonstrate. The live getDashboardData() already reads a SECOND tab (the
#   pre-May archive, located by gid) with the same column layout, so a companion
#   workbook mirrors the real architecture instead of inventing one.
#
# Sheets produced:
#   Before May 26      archive hires Jan-Apr 2026, same column layout
#   Pilot Programme         pilot membership (the live script holds this as a name list)
#   Mandatory Sessions one row per scheduled session
#   Session Attendance one row per (session, hire)
#
# Every identity here is FICTIONAL. Seeded — regenerating reproduces it exactly.
#
# The controlled vocabulary (departments, job titles, managers, recruiters,
# HRBPs, locations) is READ OUT OF THE TRACKER rather than hardcoded, so the
# archive can never introduce a parallel set of names and split the analytics
# charts in two. Only if the tracker yields nothing do the fallbacks below apply.
import random, datetime as dt
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from build_nho_data import read_hires as load_hires

random.seed(4242)
TODAY = dt.date(2026, 9, 11)          # same fixed reference date as the tracker
BASE  = os.path.dirname(os.path.abspath(__file__))   # this file's own folder
# Canonical workbook now ships one folder up (see ../_extras/WHY_THIS_IS_HERE.md);
# fall back to a local copy so this still works if that ever changes.
_LOCAL = BASE + "/NHO_Tracker_Dummy_Data.xlsx"
_PARENT = os.path.dirname(BASE) + "/NHO_Tracker_Dummy_Data.xlsx"
MAIN   = _LOCAL if os.path.exists(_LOCAL) else _PARENT
TAB   = "Cohort Tracker May - December 26"
OUT   = BASE + "/NHO_Dummy_Extensions.xlsx"

HEADERS = ["CTA Status","Manager Email Sent","Slack","New Hire Kit Ordered",
           "First Name","Last Name","Joining Date","Next Cohort Date","Personal Email",
           "Work Email","Job Title","Manager Name","Manager Email","HRBP","Notes",
           "Google Group Created","Welcome Call Invitation Sent?","Location",
           "Dashboard Link","DM to manager","DM to new hire","HRBP from sheet",
           "Department","Okta Activated","Device Ship Date"]

# ------------------------------------------------- read the tracker's own rows
tracker = load_hires(MAIN, TAB, "tracker")
if not tracker:
    raise SystemExit("No hires found in %s — nothing to extend." % MAIN)

def distinct(field, fallback):
    vals = sorted({(x.get(field) or "").strip() for x in tracker if (x.get(field) or "").strip()})
    return vals or fallback

DEPARTMENTS = distinct("department", ["Engineering", "Sales", "People"])
MANAGERS    = distinct("managerName", ["Morgan Smith", "Jamie Wilson"])
RECRUITERS  = distinct("recruiter", ["Recruiter A", "Recruiter B"])
HRBPS       = distinct("hrbp", ["HRBP A", "HRBP B"])
LOCATIONS   = distinct("location", ["Austin, Texas, United States"])

# Job titles, grouped by the department they were seen against, so an archive
# hire never ends up as a "Corporate Counsel" in Engineering.
TITLES = {}
for x in tracker:
    dept, title = (x.get("department") or "").strip(), (x.get("jobTitle") or "").strip()
    if dept and title:
        TITLES.setdefault(dept, [])
        if title not in TITLES[dept]:
            TITLES[dept].append(title)

# Manager name -> email, reusing the tracker's own address convention.
MGR_EMAIL = {}
for x in tracker:
    if x.get("managerName") and x.get("managerEmail"):
        MGR_EMAIL.setdefault(x["managerName"].strip(), x["managerEmail"].strip())

def email_domains():
    """Mirror whatever domains the tracker uses for work / personal addresses."""
    work = next((x["workEmail"].split("@")[1] for x in tracker if "@" in (x.get("workEmail") or "")), "democo.example")
    personal = next((x["personalEmail"].split("@")[1] for x in tracker if "@" in (x.get("personalEmail") or "")), "example.com")
    return work, personal
WORK_DOMAIN, PERSONAL_DOMAIN = email_domains()

# ---------------------------------------------------------- archive hire rows
# Biweekly-Monday cadence continued backwards from the tracker's first cohort.
ARCHIVE_PLAN = [("2026-01-05", 6), ("2026-01-19", 5),
                ("2026-02-02", 3), ("2026-02-16", 2),
                ("2026-03-02", 3), ("2026-03-16", 3), ("2026-03-30", 2),
                ("2026-04-13", 1)]

# 25 further fictional identities, disjoint from those in the tracker.
ARCHIVE_NAMES = [
 ("Odalys","Pemberton"),("Thaddeus","Nkemelu"),("Sunniva","Braithwaite"),
 ("Caspian","Villanueva"),("Perpetua","Oyelaran"),("Lorcan","Fitzwilliam"),
 ("Marisol","Quintanilla"),("Ignatius","Halvorsen"),("Xiomara","Bellweather"),
 ("Fionnuala","Adeyemi"),("Bartholomew","Szymanski"),("Anneliese","Karamanlis"),
 ("Cormac","Underhill"),("Solveig","Matsumura"),("Emeka","Thornbury"),
 ("Rosalind","Vukovic"),("Ptolemy","Ashgrove"),("Beatrix","Olawale"),
 ("Leander","Nakagawa"),("Verity","Castellanos"),("Osman","Lindqvist"),
 ("Clementine","Achebe"),("Rafferty","Sandoval"),("Isolde","Merriweather"),
 ("Augustin","Batbayar"),
]
tracker_names = {(x["first"], x["last"]) for x in tracker}
ARCHIVE_NAMES = [n for n in ARCHIVE_NAMES if n not in tracker_names]

def slug(s):
    return "".join(ch for ch in str(s).lower() if ch.isalnum())

archive_rows, ni = [], 0
for cohort, n in ARCHIVE_PLAN:
    cdate = dt.date.fromisoformat(cohort)
    for _ in range(n):
        if ni >= len(ARCHIVE_NAMES):
            break
        first, last = ARCHIVE_NAMES[ni]; ni += 1
        dept  = DEPARTMENTS[ni % len(DEPARTMENTS)]
        title = random.choice(TITLES.get(dept) or ["Specialist"])
        hrbp  = random.choice(HRBPS)
        mgr   = random.choice(MANAGERS)
        # Archive cohorts are entirely in the past: every onboarding step closed.
        archive_rows.append([
            "Onboarded", True, True, True,
            first, last,
            dt.datetime(cdate.year, cdate.month, cdate.day),   # Joining Date
            dt.datetime(cdate.year, cdate.month, cdate.day),   # Next Cohort Date
            "%s.%s@%s" % (slug(first), slug(last), PERSONAL_DOMAIN),
            "%s.%s@%s" % (slug(first), slug(last), WORK_DOMAIN),
            title, mgr,
            MGR_EMAIL.get(mgr, "%s@%s" % (slug(mgr), WORK_DOMAIN)),
            hrbp,
            "Dept: %s | Recruiter: %s" % (dept, random.choice(RECRUITERS)),
            "https://groups.google.com/a/%s/g/onboarding-%s" % (WORK_DOMAIN, cohort),
            True,
            random.choice(LOCATIONS),
            None, True, True, hrbp, dept,
            "✓ " + (cdate - dt.timedelta(days=1)).strftime("%Y-%m-%d"),
            dt.datetime.combine(cdate - dt.timedelta(days=5), dt.time()),
        ])

def as_date(v):
    return dt.date.fromisoformat(v[:10]) if v else None

everyone = [{"first": r[4], "last": r[5], "email": r[9],
             "start": r[6].date(), "cohort": r[7].date()} for r in archive_rows]
for x in tracker:
    start  = as_date(x["startDate"])
    cohort = as_date(x["cohortDate"]) or start
    if x["workEmail"] and start and cohort:
        everyone.append({"first": x["first"], "last": x["last"],
                         "email": x["workEmail"], "start": start, "cohort": cohort})

# ------------------------------------------------------------ Pilot pilot rows
# Mirrors the live getPilotNames_(), which is a flat name list the dashboard
# intersects with the tracker. All members have already started, so the pilot
# reads "20 members / 20 joined / 0 upcoming".
joined = [p for p in everyone if p["start"] <= TODAY]
pilot_members = sorted(random.sample(joined, min(20, len(joined))), key=lambda p: p["start"])
Pilot_TRACKS  = ["Coaching Foundations", "Manager Enablement", "Peer Circles"]
Pilot_COACHES = ["Coach Alderton", "Coach Bequette", "Coach Calloway"]
pilot_rows = []
for i, p in enumerate(pilot_members):
    pilot_rows.append([
        "%s %s" % (p["first"], p["last"]), p["email"],
        Pilot_TRACKS[i % len(Pilot_TRACKS)],
        (p["start"] + dt.timedelta(days=7)).strftime("%Y-%m-%d"),
        random.choice(["Active", "Active", "Active", "Completed"]),
        Pilot_COACHES[i % len(Pilot_COACHES)],
    ])

# ------------------------------------------------------ mandatory session rows
# offset = days after the cohort start date.
CATALOGUE = [
    ("ONB-D1",  "Day 1 — Virtual Orientation",  0,  90, "Mandatory",   "Facilitator A"),
    ("ONB-D2",  "Day 2 — How We Work",          1,  60, "Mandatory",   "Facilitator B"),
    ("ONB-BEN", "Benefits & Total Rewards",          2,  45, "Mandatory",   "Facilitator C"),
    ("ONB-SEC", "Security & Data Essentials",        3,  45, "Mandatory",   "Facilitator D"),
    ("ONB-PRD", "Product Foundations",               7,  60, "Recommended", "Facilitator E"),
    ("ONB-MGR", "Manager Essentials",               10,  90, "Recommended", "Facilitator B"),
]

cohorts = {}
for p in everyone:
    cohorts.setdefault(p["cohort"], []).append(p)

session_rows, attendance_rows = [], []
for cdate in sorted(cohorts):
    members = cohorts[cdate]
    for code, name, offset, mins, requirement, fac in CATALOGUE:
        sdate = cdate + dt.timedelta(days=offset)
        sid   = "%s-%s" % (code, sdate.strftime("%Y%m%d"))
        past  = sdate < TODAY
        session_rows.append([
            sid, name, cdate.strftime("%Y-%m-%d"), sdate.strftime("%Y-%m-%d"),
            "09:00" if offset % 2 == 0 else "14:00", mins, requirement, fac,
            len(members), "Past" if past else "Upcoming",
            "https://meet.example.com/session-" + sid.lower(),
        ])
        for p in members:
            if past:
                r = random.random()
                status = "Attended" if r < 0.84 else ("Excused" if r < 0.93 else "Absent")
            else:
                status = "Registered" if random.random() < 0.8 else "Not registered"
            attendance_rows.append([sid, name, cdate.strftime("%Y-%m-%d"), p["email"],
                                    "%s %s" % (p["first"], p["last"]), status])

# ------------------------------------------------------------------ write out
wb = Workbook()
wb.remove(wb.active)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1D1D1F")

def sheet(title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

sheet("Before May 26", HEADERS, archive_rows,
      [14,20,9,13,13,14,13,16,30,32,32,18,32,16,46,52,26,34,16,15,15,16,22,16,16])
sheet("Pilot Programme",
      ["Member Name","Work Email","Track","Enrolled On","Pilot Status","Coach"],
      pilot_rows, [24,34,24,14,14,18])
sheet("Mandatory Sessions",
      ["Session ID","Session Name","Cohort","Date","Time","Duration (min)",
       "Requirement","Facilitator","Invited","Status","Join Link"],
      session_rows, [22,32,14,12,8,14,14,16,10,11,44])
sheet("Session Attendance",
      ["Session ID","Session Name","Cohort","Work Email","Attendee","Attendance"],
      attendance_rows, [22,32,14,34,24,16])

wb.save(OUT)
print("vocabulary from tracker: %d depts, %d managers, %d recruiters, %d HRBPs, %d locations"
      % (len(DEPARTMENTS), len(MANAGERS), len(RECRUITERS), len(HRBPS), len(LOCATIONS)))
print("archive hires:   %d" % len(archive_rows))
print("pilot members:    %d" % len(pilot_rows))
print("sessions:        %d" % len(session_rows))
print("attendance rows: %d" % len(attendance_rows))
print("-> %s" % OUT)
